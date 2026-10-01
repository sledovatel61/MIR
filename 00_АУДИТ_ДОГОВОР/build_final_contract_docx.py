#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Сборка печатной версии итогового договора (DOCX) из markdown-источника.
Источник: ИТОГОВЫЙ_ДОГОВОР_ПЕРЕПЕЧАЙ_v3_FINAL.md
Результат: ИТОГОВЫЙ_ДОГОВОР_ПЕРЕПЕЧАЙ_v3_FINAL.docx
Разметка источника:
  # ...                 — заголовок документа (по центру, полужирный)
  {{CENTER:...}}        — абзац по центру
  {{CITYDATE:A|B}}      — строка «город … дата» с правым табулятором
  ## ...                — заголовок раздела
  - ...                 — элемент списка
  <N>. ...              — нумерованный элемент списка
  <!--SIG--> … === … <!--/SIG-->  — блок подписей в две колонки
  **текст**            — полужирный фрагмент
"""
import io
import os
import re

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

SRC = os.path.join(os.path.dirname(os.path.abspath(__file__)),
                   "ИТОГОВЫЙ_ДОГОВОР_ПЕРЕПЕЧАЙ_v3_FINAL.md")
DST = os.path.splitext(SRC)[0] + ".docx"

FONT = "Times New Roman"


def set_run_font(run, size=12, bold=None):
    run.font.name = FONT
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rfonts.set(qn(attr), FONT)


def add_runs(paragraph, text, size=12):
    """Разбор **полужирного** текста на runs."""
    parts = re.split(r"\*\*", text)
    for i, part in enumerate(parts):
        if not part:
            continue
        run = paragraph.add_run(part)
        set_run_font(run, size=size, bold=(i % 2 == 1))


def body_par(doc, text, indent=True, align=WD_ALIGN_PARAGRAPH.JUSTIFY):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.alignment = align
    pf.line_spacing = 1.15
    pf.space_after = Pt(4)
    if indent:
        pf.first_line_indent = Cm(1.25)
    add_runs(p, text)
    return p


def list_par(doc, text):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    pf.line_spacing = 1.15
    pf.space_after = Pt(2)
    pf.left_indent = Cm(1.1)
    pf.first_line_indent = Cm(-0.4)
    add_runs(p, text)
    return p


def heading_par(doc, text):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.alignment = WD_ALIGN_PARAGRAPH.LEFT
    pf.line_spacing = 1.15
    pf.space_before = Pt(14)
    pf.space_after = Pt(6)
    pf.keep_with_next = True
    run = p.add_run(text)
    set_run_font(run, size=12, bold=True)
    return p


def title_par(doc, text, size=14, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER,
              space_after=6):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.alignment = align
    pf.space_after = Pt(space_after)
    pf.line_spacing = 1.15
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold)
    return p


def citydate_par(doc, left, right, text_width_cm=16.5):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.alignment = WD_ALIGN_PARAGRAPH.LEFT
    pf.space_after = Pt(10)
    pf.line_spacing = 1.15
    pf.tab_stops.add_tab_stop(Cm(text_width_cm), WD_TAB_ALIGNMENT.RIGHT)
    run = p.add_run(left + "\t" + right)
    set_run_font(run, size=12)
    return p


def no_borders(table):
    tbl_pr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement("w:" + edge)
        el.set(qn("w:val"), "none")
        el.set(qn("w:sz"), "0")
        borders.append(el)
    tbl_pr.append(borders)


def cant_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:cantSplit")
    tr_pr.append(el)


def add_page_numbers(section):
    footer = section.footer
    p = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("стр. ")
    set_run_font(run, size=10)
    run2 = p.add_run()
    set_run_font(run2, size=10)
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run2._r.append(fld_begin)
    run2._r.append(instr)
    run2._r.append(fld_end)


def main():
    text = io.open(SRC, encoding="utf-8").read()
    lines = text.split("\n")

    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = FONT
    st.font.size = Pt(12)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)

    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    sec.left_margin, sec.right_margin = Cm(3.0), Cm(1.5)
    sec.top_margin, sec.bottom_margin = Cm(2.0), Cm(2.0)
    add_page_numbers(sec)

    i = 0
    while i < len(lines):
        line = lines[i].rstrip()
        if not line.strip():
            i += 1
            continue
        if line.startswith("# "):
            title_par(doc, line[2:].strip(), size=14)
        elif line.startswith("{{CENTER:"):
            m = re.match(r"\{\{CENTER:(.*)\}\}$", line)
            title_par(doc, m.group(1), size=12, bold=False, space_after=6)
        elif line.startswith("{{CITYDATE:"):
            m = re.match(r"\{\{CITYDATE:(.*)\|(.*)\}\}$", line)
            citydate_par(doc, m.group(1), m.group(2))
        elif line.startswith("<!--SIG-->"):
            block = []
            i += 1
            while i < len(lines) and not lines[i].startswith("<!--/SIG-->"):
                block.append(lines[i].rstrip())
                i += 1
            left, right = [], []
            target = left
            for ln in block:
                if ln.strip() == "===":
                    target = right
                    continue
                target.append(ln)
            table = doc.add_table(rows=1, cols=2)
            no_borders(table)
            cant_split(table.rows[0])
            for cell, chunk in zip(table.rows[0].cells, (left, right)):
                cell.paragraphs[0]._element.getparent().remove(
                    cell.paragraphs[0]._element)
                for ln in chunk:
                    p = cell.add_paragraph()
                    pf = p.paragraph_format
                    pf.space_after = Pt(2)
                    pf.line_spacing = 1.15
                    add_runs(p, ln.strip(), size=12)
        elif line.startswith("## "):
            heading_par(doc, line[3:].strip())
        elif line.startswith("- "):
            list_par(doc, "– " + line[2:].strip())
        elif re.match(r"^\d+\.\s", line):
            list_par(doc, line.strip())
        else:
            body_par(doc, line.strip())
        i += 1

    doc.save(DST)
    print("saved:", DST, os.path.getsize(DST), "bytes")


if __name__ == "__main__":
    main()
