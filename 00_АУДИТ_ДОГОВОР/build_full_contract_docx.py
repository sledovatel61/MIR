#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Сборка единого печатного документа: Договор № 12-ОУТ + Приложения № 1–10.
Источники:
  ИТОГОВЫЙ_ДОГОВОР_ПЕРЕПЕЧАЙ_v3_FINAL.md
  ПРИЛОЖЕНИЯ_К_ДОГОВОРУ_12-ОУТ.md
Результат:
  ДОГОВОР_12-ОУТ_С_ПРИЛОЖЕНИЯМИ.docx

Поддерживается: заголовки, **полужирный**, *курсив*, списки, таблицы (| |),
блок-цитаты (>), разметка {{CENTER:}}, {{CITYDATE:}}, <!--SIG-->, разрывы страниц
между приложениями, альбомная ориентация для широких таблиц.
"""
import io
import os
import re

from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

HERE = os.path.dirname(os.path.abspath(__file__))
CONTRACT_MD = os.path.join(HERE, "ИТОГОВЫЙ_ДОГОВОР_ПЕРЕПЕЧАЙ_v3_FINAL.md")
APPENDIX_MD = os.path.join(HERE, "ПРИЛОЖЕНИЯ_К_ДОГОВОРУ_12-ОУТ.md")
DST = os.path.join(HERE, "ДОГОВОР_12-ОУТ_С_ПРИЛОЖЕНИЯМИ.docx")

FONT = "Times New Roman"
PORTRAIT_W, PORTRAIT_H = Cm(21.0), Cm(29.7)
LAND_W, LAND_H = Cm(29.7), Cm(21.0)


# ------------------------- базовые помощники -------------------------

def set_run_font(run, size=12, bold=None, italic=None):
    run.font.name = FONT
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rfonts.set(qn(attr), FONT)


def add_runs(paragraph, text, size=12, base_bold=False):
    """**жирный**, *курсив*."""
    for part in re.split(r"(\*\*.+?\*\*)", text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**") and len(part) > 4:
            run = paragraph.add_run(part[2:-2])
            set_run_font(run, size=size, bold=True)
            continue
        for sub in re.split(r"(\*[^*\n]+\*)", part):
            if not sub:
                continue
            if sub.startswith("*") and sub.endswith("*") and len(sub) > 2:
                run = paragraph.add_run(sub[1:-1])
                set_run_font(run, size=size, italic=True, bold=base_bold or None)
            else:
                run = paragraph.add_run(sub)
                set_run_font(run, size=size, bold=True if base_bold else None)


def body_par(doc, text, size=12, indent=True, align=WD_ALIGN_PARAGRAPH.JUSTIFY,
             space_after=4, left_indent=None):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.alignment = align
    pf.line_spacing = 1.15
    pf.space_after = Pt(space_after)
    if indent:
        pf.first_line_indent = Cm(1.25)
    if left_indent is not None:
        pf.left_indent = Cm(left_indent)
    add_runs(p, text, size=size)
    return p


def list_par(doc, text, size=12):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    pf.line_spacing = 1.15
    pf.space_after = Pt(2)
    pf.left_indent = Cm(1.1)
    pf.first_line_indent = Cm(-0.4)
    add_runs(p, text, size=size)
    return p


def quote_par(doc, text, size=11):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    pf.line_spacing = 1.1
    pf.space_after = Pt(2)
    pf.left_indent = Cm(0.75)
    pf.right_indent = Cm(0.5)
    add_runs(p, text, size=size, base_bold=False)
    return p


def heading_par(doc, text, size=12):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.alignment = WD_ALIGN_PARAGRAPH.LEFT
    pf.line_spacing = 1.15
    pf.space_before = Pt(12)
    pf.space_after = Pt(6)
    pf.keep_with_next = True
    run = p.add_run(text)
    set_run_font(run, size=size, bold=True)
    return p


def title_par(doc, text, size=14, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER,
              space_after=6, space_before=0):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.alignment = align
    pf.space_after = Pt(space_after)
    pf.space_before = Pt(space_before)
    pf.line_spacing = 1.15
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold)
    return p


def citydate_par(doc, left, right, text_width_cm):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.alignment = WD_ALIGN_PARAGRAPH.LEFT
    pf.space_after = Pt(10)
    pf.line_spacing = 1.15
    pf.tab_stops.add_tab_stop(Cm(text_width_cm), WD_TAB_ALIGNMENT.RIGHT)
    run = p.add_run(left + "\t" + right)
    set_run_font(run, size=12)
    return p


def no_borders(table):
    tbl_pr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement("w:" + edge)
        el.set(qn("w:val"), "none")
        el.set(qn("w:sz"), "0")
        borders.append(el)
    tbl_pr.append(borders)


def cant_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    tr_pr.append(OxmlElement("w:cantSplit"))


def repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:tblHeader")
    el.set(qn("w:val"), "true")
    tr_pr.append(el)


def add_page_numbers(section):
    footer = section.footer
    footer.is_linked_to_previous = False
    p = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for r in list(p.runs):
        r._element.getparent().remove(r._element)
    run = p.add_run("стр. ")
    set_run_font(run, size=10)
    run2 = p.add_run()
    set_run_font(run2, size=10)
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run2._r.append(fld_begin)
    run2._r.append(instr)
    run2._r.append(fld_end)


def setup_section(section, landscape=False):
    if landscape:
        section.orientation = WD_ORIENT.LANDSCAPE
        section.page_width, section.page_height = LAND_W, LAND_H
        section.left_margin, section.right_margin = Cm(2.0), Cm(1.5)
        section.top_margin, section.bottom_margin = Cm(1.5), Cm(1.5)
    else:
        section.orientation = WD_ORIENT.PORTRAIT
        section.page_width, section.page_height = PORTRAIT_W, PORTRAIT_H
        section.left_margin, section.right_margin = Cm(3.0), Cm(1.5)
        section.top_margin, section.bottom_margin = Cm(2.0), Cm(2.0)
    add_page_numbers(section)


def add_md_table(doc, rows, landscape=False):
    """rows — список списков ячеек (первая строка — заголовок)."""
    ncols = max(len(r) for r in rows)
    table = doc.add_table(rows=0, cols=ncols)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    size = 10 if ncols <= 4 else (9 if ncols <= 7 else 8)
    for ri, row in enumerate(rows):
        cells = table.add_row().cells
        for ci in range(ncols):
            text = row[ci] if ci < len(row) else ""
            cell = cells[ci]
            p = cell.paragraphs[0]
            pf = p.paragraph_format
            pf.space_after = Pt(0)
            pf.space_before = Pt(0)
            pf.line_spacing = 1.0
            add_runs(p, text, size=size, base_bold=(ri == 0))
        cant_split(table.rows[-1])
    if rows:
        repeat_header(table.rows[0])
    return table


# ------------------------- парсинг разметки -------------------------

def parse_table(lines, i):
    rows = []
    while i < len(lines) and lines[i].lstrip().startswith("|"):
        raw = lines[i].strip().strip("|")
        cells = [c.strip() for c in raw.split("|")]
        if not all(re.fullmatch(r":?-{2,}:?", c or "-") for c in cells):
            rows.append(cells)
        i += 1
    return rows, i


def render_contract(doc, text):
    lines = text.split("\n")
    i = 0
    while i < len(lines):
        line = lines[i].rstrip()
        s = line.strip()
        if not s:
            i += 1
            continue
        if s.startswith("<!--SIG-->"):
            block = []
            i += 1
            while i < len(lines) and not lines[i].startswith("<!--/SIG-->"):
                block.append(lines[i].rstrip())
                i += 1
            left, right, target = [], [], None
            for ln in block:
                if ln.strip() == "===":
                    target = right
                    continue
                (left if target is None else right).append(ln)
            table = doc.add_table(rows=1, cols=2)
            no_borders(table)
            cant_split(table.rows[0])
            for cell, chunk in zip(table.rows[0].cells, (left, right)):
                cell.paragraphs[0]._element.getparent().remove(
                    cell.paragraphs[0]._element)
                for ln in chunk:
                    p = cell.add_paragraph()
                    pf = p.paragraph_format
                    pf.space_after = Pt(2)
                    pf.line_spacing = 1.15
                    add_runs(p, ln.strip(), size=12)
        elif s.startswith("|"):
            rows, i = parse_table(lines, i)
            if rows:
                add_md_table(doc, rows)
            continue
        elif line.startswith("# "):
            title_par(doc, line[2:].strip(), size=14)
        elif line.startswith("## "):
            heading_par(doc, line[3:].strip())
        elif line.startswith("{{CENTER:"):
            m = re.match(r"\{\{CENTER:(.*)\}\}$", s)
            title_par(doc, m.group(1), size=12, bold=False, space_after=6)
        elif line.startswith("{{CITYDATE:"):
            m = re.match(r"\{\{CITYDATE:(.*)\|(.*)\}\}$", s)
            citydate_par(doc, m.group(1), m.group(2), 16.5)
        elif s.startswith("- "):
            list_par(doc, "– " + s[2:].strip())
        elif re.match(r"^\d+[\.\)]\s", s):
            list_par(doc, s)
        else:
            body_par(doc, s)
        i += 1


APPENDIX_RE = re.compile(r"^##\s+\d+\.\s+ПРИЛОЖЕНИЕ\s+№\s+(\d+)\s+к\s+Договору\s+№\s+12-ОУТ\s+—\s+(.+)$")


def render_appendices(doc):
    text = io.open(APPENDIX_MD, encoding="utf-8").read()
    lines = text.split("\n")
    sections = []
    cur_title, cur_num, cur_lines, started = None, None, [], False
    for ln in lines:
        m = APPENDIX_RE.match(ln.strip())
        if m:
            if started:
                sections.append((cur_num, cur_title, cur_lines))
            cur_num, cur_title, cur_lines, started = m.group(1), m.group(2), [], True
            continue
        if started:
            cur_lines.append(ln)
    if started:
        sections.append((cur_num, cur_title, cur_lines))

    assert len(sections) == 10, f"Приложений найдено: {len(sections)}"

    for num, title, body in sections:
        # широкая таблица => альбомная ориентация
        maxcols = 0
        for ln in body:
            if ln.strip().startswith("|"):
                maxcols = max(maxcols, ln.strip().strip("|").count("|") + 1)
        # принудительная ориентация: реестр (№ 2) — книжная, график (№ 4) — альбомная
        ORIENT_OVERRIDE = {"2": False, "4": True}
        landscape = ORIENT_OVERRIDE.get(num, maxcols >= 8)
        sec = doc.add_section(WD_SECTION.NEW_PAGE)
        setup_section(sec, landscape=landscape)

        title_par(doc, f"ПРИЛОЖЕНИЕ № {num}", size=14, space_after=2,
                  space_before=0)
        title_par(doc, "к Договору № 12-ОУТ на оказание услуг по кадровому аутсорсингу",
                  size=11, bold=False, space_after=2)
        title_par(doc, title.strip(), size=12, space_after=10)
        if num == "4":
            body = [ln for ln in body]
        i = 0
        while i < len(body):
            line = body[i].rstrip()
            s = line.strip()
            if not s or s == "---":
                i += 1
                continue
            if s.startswith("|"):
                rows, i = parse_table(body, i)
                if rows:
                    add_md_table(doc, rows, landscape=landscape)
                    doc.add_paragraph().paragraph_format.space_after = Pt(4)
                continue
            if s.startswith(">"):
                quote_par(doc, s.lstrip("> ").strip())
            elif s.startswith("**") and s.endswith("**") and len(s) < 160:
                p = doc.add_paragraph()
                pf = p.paragraph_format
                pf.space_before = Pt(8)
                pf.space_after = Pt(4)
                pf.keep_with_next = True
                pf.line_spacing = 1.15
                add_runs(p, s, size=12)
            elif s.startswith("- "):
                list_par(doc, "– " + s[2:].strip())
            elif re.match(r"^\d+[\.\)]\s", s):
                list_par(doc, s)
            else:
                body_par(doc, s)
            i += 1


def main():
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = FONT
    st.font.size = Pt(12)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    setup_section(doc.sections[0], landscape=False)

    render_contract(doc, io.open(CONTRACT_MD, encoding="utf-8").read())
    render_appendices(doc)

    doc.save(DST)
    print("saved:", DST, os.path.getsize(DST), "bytes")


if __name__ == "__main__":
    main()
