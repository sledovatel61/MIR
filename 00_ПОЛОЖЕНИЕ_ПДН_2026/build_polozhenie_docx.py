#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Сборка печатной (готовой) версии Положения об обработке ПДн ООО «МИР» в DOCX
из выверенного markdown-источника v1.0.

Разметка источника:
  # ...                    — часть/крупный раздел (с новой страницы, по центру)
  ## ...                   — заголовок 1-го уровня (приложение — с новой страницы)
  ### ...                  — заголовок 2-го уровня
  **текст**                — полужирный фрагмент;  *текст* — курсив
  > ...                    — форма/блок-врезка (с отступом); «> | ... |» — таблица формы
  | ... | ... |            — таблица
  - ... / N) ...           — списки
  ---                      — разделитель (пропускается)
"""
import io
import os
import re

from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

BASE = os.path.dirname(os.path.abspath(__file__))
LANDSCAPE_APPS = {11, 12, 13, 14, 15, 20}   # широкие журналы и реестры — альбомная ориентация


def switch_section(doc, landscape):
    """Новый раздел с нужной ориентацией (альбомная — для широких таблиц)."""
    sec = doc.add_section(WD_SECTION.NEW_PAGE)
    if landscape:
        sec.orientation = WD_ORIENT.LANDSCAPE
        sec.page_width, sec.page_height = Cm(29.7), Cm(21.0)
        sec.top_margin = sec.bottom_margin = Cm(1.5)
        sec.left_margin, sec.right_margin = Cm(1.5), Cm(1.2)
    else:
        sec.orientation = WD_ORIENT.PORTRAIT
        sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
        sec.top_margin = sec.bottom_margin = Cm(2.0)
        sec.left_margin, sec.right_margin = Cm(2.8), Cm(1.5)
    return sec
SRC = os.path.join(BASE, "ПОЛОЖЕНИЕ_ОБ_ОБРАБОТКЕ_ПДН_ООО_МИР_v1.0.md")
DST = os.path.join(BASE, "ПОЛОЖЕНИЕ_ОБ_ОБРАБОТКЕ_ПДН_ООО_МИР_v1.0.docx")

FONT = "Times New Roman"
ACCENT = RGBColor(0x1F, 0x37, 0x63)


def set_run(run, size=12, bold=None, italic=None, color=None, font=FONT):
    run.font.name = font
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color is not None:
        run.font.color.rgb = color
    rpr = run._element.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts")
        rpr.append(rf)
    for a in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rf.set(qn(a), font)


TOKEN = re.compile(r"(\*\*[^*]+\*\*|\*[^*\n]+\*|`[^`]+`|\[[^\]\n]+\])")


def add_inline(par, text, size=12, base_bold=False):
    """Разбор **жирного**, *курсива*, `кода` и [ПЛЕЙСХОЛДЕРОВ] на runs."""
    pos = 0
    for m in TOKEN.finditer(text):
        if m.start() > pos:
            r = par.add_run(text[pos:m.start()])
            set_run(r, size=size, bold=True if base_bold else None)
        tok = m.group(0)
        if tok.startswith("**"):
            r = par.add_run(tok[2:-2]); set_run(r, size=size, bold=True)
        elif tok.startswith("`"):
            r = par.add_run(tok[1:-1]); set_run(r, size=size, font="Consolas")
        elif tok.startswith("*"):
            r = par.add_run(tok[1:-1]); set_run(r, size=size, italic=True)
        else:  # [ПЛЕЙСХОЛДЕР] — выделяем, чтобы не забыть заполнить
            r = par.add_run(tok); set_run(r, size=size, bold=True, color=RGBColor(0xB0, 0x30, 0x30))
        pos = m.end()
    if pos < len(text):
        r = par.add_run(text[pos:]); set_run(r, size=size, bold=True if base_bold else None)


def para(doc, text="", size=12, align=WD_ALIGN_PARAGRAPH.JUSTIFY, indent=1.25,
         before=0, after=6, left=0, hanging=None, space_line=1.15, keep_next=False):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.alignment = align
    pf.first_line_indent = Cm(indent) if indent else None
    pf.left_indent = Cm(left) if left else None
    if hanging:
        pf.first_line_indent = Cm(-hanging)
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = space_line
    pf.keep_with_next = keep_next
    add_inline(p, text, size=size)
    return p


def heading(doc, text, level, page_break=False):
    p = doc.add_paragraph()
    if page_break:
        p.paragraph_format.page_break_before = True
    pf = p.paragraph_format
    pf.keep_with_next = True
    pf.space_before = Pt(18 if level == 1 else 12)
    pf.space_after = Pt(10 if level == 1 else 6)
    pf.line_spacing = 1.1
    pf.alignment = WD_ALIGN_PARAGRAPH.CENTER if level == 1 else WD_ALIGN_PARAGRAPH.LEFT
    sizes = {1: 15, 2: 13, 3: 12}
    add_inline(p, text, size=sizes.get(level, 12))
    for r in p.runs:
        r.bold = True
        r.font.color.rgb = ACCENT if level <= 2 else RGBColor(0, 0, 0)
    return p


def shade(cell, fill="DCE6F1"):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tcPr.append(shd)


def repeat_header(row):
    trPr = row._tr.get_or_add_trPr()
    th = OxmlElement("w:tblHeader")
    th.set(qn("w:val"), "true")
    trPr.append(th)


def set_tbl_width(tbl, pct=5000):
    tblPr = tbl._tbl.tblPr
    w = OxmlElement("w:tblW")
    w.set(qn("w:type"), "pct")
    w.set(qn("w:w"), str(pct))
    tblPr.append(w)
    mar = OxmlElement("w:tblCellMar")
    for side, val in (("left", 40), ("right", 40), ("top", 20), ("bottom", 20)):
        el = OxmlElement("w:" + side)
        el.set(qn("w:w"), str(val))
        el.set(qn("w:type"), "dxa")
        mar.append(el)
    tblPr.append(mar)


def add_table(doc, rows, size=None, left=None):
    ncols = max(len(r) for r in rows)
    if size is None:
        size = 7.5 if ncols >= 9 else (8.5 if ncols >= 6 else 10)
    tbl = doc.add_table(rows=len(rows), cols=ncols)
    tbl.style = "Table Grid"
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = True
    set_tbl_width(tbl)
    for i, r in enumerate(rows):
        for j in range(ncols):
            cell = tbl.cell(i, j)
            txt = r[j] if j < len(r) else ""
            cell.text = ""
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.line_spacing = 1.0
            add_inline(p, txt.strip(), size=size, base_bold=(i == 0))
            if i == 0:
                shade(cell)
    repeat_header(tbl.rows[0])
    if left:
        # отступ врезки для таблиц форм — через пустой абзац не делаем, прижимаем к тексту
        pass
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return tbl


def split_row(line):
    s = line.strip()
    if s.startswith("|"):
        s = s[1:]
    if s.endswith("|"):
        s = s[:-1]
    return [c.strip() for c in s.split("|")]


def is_sep(line):
    s = line.strip()
    return bool(re.fullmatch(r"\|?[\s:\-|]+\|?", s)) and "-" in s and set(s) <= set("|-: ")


def add_footer(section, text):
    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(text + "   |   стр. ")
    set_run(r, size=9, color=RGBColor(0x60, 0x60, 0x60))
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    sz = OxmlElement("w:sz"); sz.set(qn("w:val"), "18"); rpr.append(sz)
    run.append(rpr)
    t = OxmlElement("w:t"); t.text = "1"; run.append(t)
    fld.append(run)
    p._p.append(fld)


def build():
    text = io.open(SRC, encoding="utf-8").read()
    lines = text.split("\n")

    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = FONT
    st.font.size = Pt(12)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)

    sec = doc.sections[0]
    sec.page_height = Cm(29.7)
    sec.page_width = Cm(21.0)
    sec.top_margin = Cm(2.0)
    sec.bottom_margin = Cm(2.0)
    sec.left_margin = Cm(2.8)
    sec.right_margin = Cm(1.5)
    add_footer(sec, "Положение об обработке и защите персональных данных ООО «МИР»")

    # ---------- титульный лист ----------
    p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(0)
    r = p.add_run("ОБЩЕСТВО С ОГРАНИЧЕННОЙ ОТВЕТСТВЕННОСТЬЮ «МИР»")
    set_run(r, size=13, bold=True); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("ОГРН 1156196053881 · ИНН 6166094734")
    set_run(r, size=10, color=RGBColor(0x50, 0x50, 0x50))
    r.add_break()
    r2 = p.add_run("344093, г. Ростов-на-Дону, ул. Туполева, д. 16Е, оф. 4.13")
    set_run(r2, size=10, color=RGBColor(0x50, 0x50, 0x50))
    for _ in range(6):
        doc.add_paragraph()
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run("ПОЛОЖЕНИЕ\nоб обработке и защите персональных данных работников\nи иных субъектов персональных данных")
    set_run(r, size=20, bold=True, color=ACCENT)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Выверенный проект v1.0\n(по итогам аудита трёх независимых агентов, 01.10.2026)")
    set_run(r, size=12, italic=True, color=RGBColor(0x50, 0x50, 0x50))
    for _ in range(8):
        doc.add_paragraph()
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.left_indent = Cm(7.5)
    r = p.add_run("УТВЕРЖДЕНО\nприказом Генерального директора\nООО «МИР»\nот «___» __________ 20___ г. № ___")
    set_run(r, size=12)
    for _ in range(3):
        doc.add_paragraph()
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("г. Ростов-на-Дону · 2026")
    set_run(r, size=12)

    # ---------- содержание ----------
    doc.add_page_break()
    heading(doc, "СОДЕРЖАНИЕ", 1)
    toc = [
        "Часть I. Общие положения (п. 1)",
        "Часть II. Политика обработки персональных данных: принципы, основания, цели, состав данных (п. 2–6)",
        "Часть III. Регламент обработки: сбор, доступ, хранение, передача, уничтожение, защита, инциденты (п. 7–17)",
        "Часть IV. Права субъектов, обращения, ответственность (п. 18–20)",
        "Часть V. Заключительные положения, внедрение, карта соответствия (п. 21–23)",
        "Приложение № 1. Политика в отношении обработки персональных данных (публичная редакция)",
        "Приложение № 2. Перечень должностей, замещение которых предусматривает обработку персональных данных",
        "Приложение № 3. Перечень персональных данных, обрабатываемых по целям",
        "Приложение № 4. Матрица обработки (цель → основание → срок → ответственный → место хранения)",
        "Приложение № 5. Форма согласия работника на обработку персональных данных",
        "Приложение № 6. Форма согласия кандидата на обработку персональных данных",
        "Приложение № 7. Форма согласия на обработку персональных данных, разрешённых для распространения",
        "Приложение № 8. Форма согласия на обработку персональных данных несовершеннолетних детей",
        "Приложение № 9. Обязательство о неразглашении персональных данных",
        "Приложение № 10. Типовое поручение на обработку персональных данных (для обработчиков)",
        "Приложение № 11. Реестр обработчиков персональных данных",
        "Приложение № 12. Журнал учёта машинных носителей персональных данных",
        "Приложение № 13. Журнал учёта инцидентов и формы уведомлений в Роскомнадзор",
        "Приложение № 14. Форма акта об уничтожении персональных данных",
        "Приложение № 15. Лист ознакомления и журнал учёта обращений субъектов персональных данных",
        "Приложение № 16. Чек-лист приёма на работу",
        "Приложение № 17. Чек-лист увольнения работника",
        "Приложение № 18. Чек-лист действий при инциденте",
        "Приложение № 19. Форма акта оценки вреда субъектам персональных данных",
        "Приложение № 20. Журнал учёта передачи персональных данных третьим лицам",
        "Примечания к выверенному проекту (служебный раздел: что заполнить до утверждения)",
    ]
    for i, item in enumerate(toc):
        p = doc.add_paragraph()
        pf = p.paragraph_format
        pf.space_after = Pt(2)
        pf.line_spacing = 1.15
        pf.left_indent = Cm(0.6 if i > 4 else 0)
        add_inline(p, item, size=11.5)

    # ---------- тело ----------
    i = 0
    n = len(lines)
    landscape_now = False
    while i < n:
        raw = lines[i]
        line = raw.rstrip()
        s = line.strip()
        if not s or s == "---":
            i += 1
            continue

        # заголовки
        if s.startswith("# "):
            title = s[2:].strip()
            if title.startswith("ПОЛОЖЕНИЕ об обработке"):
                i += 1
                continue  # титул уже собран
            if title.startswith("ПРИМЕЧАНИЯ"):
                if landscape_now:
                    switch_section(doc, False)
                    landscape_now = False
                heading(doc, "ПРИМЕЧАНИЯ К ВЫВЕРЕННОМУ ПРОЕКТУ", 1, page_break=True)
            elif title.startswith("ПРИЛОЖЕНИЯ"):
                heading(doc, "ПРИЛОЖЕНИЯ", 1, page_break=True)
            else:
                heading(doc, title, 1, page_break=True)
            i += 1
            continue
        if s.startswith("## "):
            title = s[3:].strip()
            mapp = re.match(r"Приложение № (\d+)\.", title)
            want_land = bool(mapp) and int(mapp.group(1)) in LANDSCAPE_APPS
            if want_land != landscape_now:
                switch_section(doc, want_land)
                landscape_now = want_land
                pb = False
            else:
                pb = title.startswith("Приложение №") or title.startswith("ПРИМЕЧАНИЯ")
            heading(doc, title, 2, page_break=pb)
            i += 1
            continue
        if s.startswith("### "):
            heading(doc, s[4:].strip(), 3)
            i += 1
            continue

        # врезки-формы
        if s.startswith(">"):
            block = []
            while i < n and lines[i].strip().startswith(">"):
                block.append(lines[i].strip()[1:].strip())
                i += 1
            j = 0
            while j < len(block):
                b = block[j]
                if b.startswith("|"):
                    rows = []
                    while j < len(block) and block[j].startswith("|"):
                        if not is_sep(block[j]):
                            rows.append(split_row(block[j]))
                        j += 1
                    if rows:
                        add_table(doc, rows, left=True)
                    continue
                if b == "":
                    j += 1
                    continue
                para(doc, b, size=11.5, indent=0, left=0.8, after=4)
                j += 1
            continue

        # таблицы
        if s.startswith("|"):
            rows = []
            while i < n and lines[i].strip().startswith("|"):
                if not is_sep(lines[i]):
                    rows.append(split_row(lines[i]))
                i += 1
            if rows:
                add_table(doc, rows)
            continue

        # списки
        m = re.match(r"^(\d+)\)\s+(.*)$", s)
        if m:
            p = doc.add_paragraph()
            pf = p.paragraph_format
            pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            pf.left_indent = Cm(1.1)
            pf.first_line_indent = Cm(-0.55)
            pf.space_after = Pt(3)
            pf.line_spacing = 1.15
            add_inline(p, m.group(1) + ") " + m.group(2), size=12)
            i += 1
            continue
        if s.startswith("- "):
            para(doc, "• " + s[2:], indent=0, left=0.8, after=3)
            i += 1
            continue
        if s.startswith("*") and s.endswith("*") and not s.startswith("**"):
            para(doc, s, size=11, align=WD_ALIGN_PARAGRAPH.LEFT, indent=0, after=6)
            i += 1
            continue

        para(doc, s)
        i += 1

    doc.save(DST)
    print("Сохранено:", DST)


if __name__ == "__main__":
    build()
